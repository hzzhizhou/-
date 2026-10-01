"""表格文档解析：CSV / DOCX / XLSX 的表格统一转成「带表头的 Markdown 表格」文本块。

为什么不用现成 loader：
- Docx2txtLoader 把表格退化成单元格文字顺序拼接，行列对应关系丢失；
- UnstructuredExcelLoader 输出逐格 "column: value" 文本，同样丢失表格结构。
表格一旦被当成普通段落文本，下游分块会按标点/换行把它切散：表头与数据行分离、
一行里的多个列被拆进不同的块，检索命中 "599" 这类单元格时无法还原它属于哪个商品。

分块约定（与分块层的 is_table 保护配套）：
- 每块 = 表头 + 分隔行 + 若干条完整数据行；块内绝不切开某一行，也绝不让数据行脱离表头；
- 每块输出一个 Document，metadata 标 is_table="true"，分块层据此跳过二次切分。
"""
import csv
from pathlib import Path
from typing import List, Sequence

from langchain_core.documents import Document

from config.settings import TABLE_CHUNK_SIZE
from logs.log_config import data_layer_log as log

# 单块字符预算：走独立的 TABLE_CHUNK_SIZE，不沿用子块大小。表格每块都要重复表头，
# 宽表的表头本身就很长（9 列约 144 字符），借用 400 的小预算会把整表碎成
# 大量「表头 + 一行」的块（实测 40 行 FAQ 切成 28 块）。
DEFAULT_TABLE_BUDGET = TABLE_CHUNK_SIZE


def _cell(value) -> str:
    """单元格文本：压平换行、去掉竖线（竖线会破坏 Markdown 表格结构）"""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)          # xlsx 里的 599 会读成 599.0
    text = str(value).replace("\r", " ").replace("\n", " ")
    return " ".join(text.split()).replace("|", "/")


def render_markdown_table(header: Sequence[str], rows: Sequence[Sequence[str]]) -> str:
    """渲染 Markdown 表格（含表头分隔行）"""
    head = "| " + " | ".join(header) + " |"
    separator = "| " + " | ".join(["---"] * len(header)) + " |"
    body = ["| " + " | ".join(row) + " |" for row in rows]
    return "\n".join([head, separator] + body)


def _normalize_rows(header: Sequence[str], rows: List[List[str]]) -> List[List[str]]:
    """列数对齐表头：短行补空、长行截断，避免 Markdown 表格列错位"""
    width = len(header)
    fixed = []
    for row in rows:
        if len(row) < width:
            fixed.append(list(row) + [""] * (width - len(row)))
        elif len(row) > width:
            fixed.append(list(row[:width]))
        else:
            fixed.append(list(row))
    return fixed


def _table_documents(header: List[str], rows: List[List[str]],
                     metadata: dict, budget: int) -> List[Document]:
    """
    按字符预算把整表切成多个「表头 + 整行」块
    :param header: 表头单元格
    :param rows: 数据行（已对齐列数）
    :param metadata: 附加元数据（除 is_table 外）
    :return: 表格块列表（无数据行时返回空，避免生成只有列名的无效块）
    """
    if not header or not rows:
        return []
    docs = []
    prefix_len = len(render_markdown_table(header, []))          # 表头 + 分隔行
    group: List[List[str]] = []
    size = prefix_len
    groups: List[List[List[str]]] = []
    for row in rows:
        row_len = len("| " + " | ".join(row) + " |") + 1
        if group and size + row_len > budget:
            groups.append(group)
            group, size = [], prefix_len
        group.append(row)
        size += row_len
    if group:
        groups.append(group)

    for part in groups:
        meta = dict(metadata)
        meta["is_table"] = "true"
        meta["table_rows"] = len(part)
        docs.append(Document(page_content=render_markdown_table(header, part), metadata=meta))
    return docs


def csv_to_documents(file_path: Path, budget: int = DEFAULT_TABLE_BUDGET) -> List[Document]:
    """CSV → 带表头的 Markdown 表格块（首行视为表头）"""
    with open(file_path, "r", encoding="utf-8-sig", newline="") as f:
        raw_rows = [[_cell(c) for c in row] for row in csv.reader(f)]
    raw_rows = [row for row in raw_rows if any(row)]
    if len(raw_rows) < 2:
        log.warning(f"CSV 无有效数据行，跳过: {file_path.name}")
        return []
    header = raw_rows[0]
    return _table_documents(header, _normalize_rows(header, raw_rows[1:]), {}, budget)


def docx_to_documents(file_path: Path, budget: int = DEFAULT_TABLE_BUDGET) -> List[Document]:
    """DOCX → 段落文本块 + 表格 Markdown 块（按正文中出现顺序输出）"""
    from docx import Document as DocxFile
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    docx = DocxFile(str(file_path))
    docs: List[Document] = []
    buffer: List[str] = []
    table_index = 0

    for child in docx.element.body.iterchildren():
        if child.tag == qn("w:p"):
            text = Paragraph(child, docx).text.strip()
            if text:
                buffer.append(text)
        elif child.tag == qn("w:tbl"):
            if buffer:
                docs.append(Document(page_content="\n".join(buffer),
                                     metadata={"is_table": "false"}))
                buffer = []
            table_rows = [[_cell(cell.text) for cell in row.cells]
                          for row in Table(child, docx).rows]
            table_rows = [row for row in table_rows if any(row)]
            if len(table_rows) < 2:
                # 单行表（多为排版用表）：按普通文本保内容，不足两行无法构成表头+数据
                if table_rows:
                    docs.append(Document(page_content=" | ".join(table_rows[0]),
                                         metadata={"is_table": "false"}))
                continue
            header = table_rows[0]
            docs.extend(_table_documents(
                header, _normalize_rows(header, table_rows[1:]),
                {"table_index": table_index}, budget))
            table_index += 1

    if buffer:
        docs.append(Document(page_content="\n".join(buffer), metadata={"is_table": "false"}))
    return docs


def xlsx_to_documents(file_path: Path, budget: int = DEFAULT_TABLE_BUDGET) -> List[Document]:
    """XLSX → 每个工作表一张 Markdown 表（首行视为表头，按预算分成多块）"""
    from openpyxl import load_workbook

    workbook = load_workbook(str(file_path), data_only=True, read_only=True)
    docs: List[Document] = []
    try:
        for sheet in workbook.worksheets:
            raw_rows = [[_cell(c) for c in row] for row in sheet.iter_rows(values_only=True)]
            raw_rows = [row for row in raw_rows if any(row)]
            if len(raw_rows) < 2:
                log.warning(f"工作表 {sheet.title} 无有效数据行，跳过: {file_path.name}")
                continue
            header = raw_rows[0]
            docs.extend(_table_documents(
                header, _normalize_rows(header, raw_rows[1:]),
                {"sheet_name": str(sheet.title)}, budget))
    finally:
        workbook.close()
    return docs